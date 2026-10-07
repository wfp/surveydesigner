import importlib
from io import BytesIO

from core.utils import get_model_admin_base_url
from django.apps import apps as django_apps
from django.forms import modelform_factory
from django.urls import reverse
from openpyxl import load_workbook
from organization.models import Organization
from questions.forms import (
    CalculationAdminModelForm,
    RepeatSectionAdminModelForm,
    SubQuestionAdminModelForm,
)
from questions.models import (
    BaseQuestion,
    Calculation,
    RepeatSection,
    RootQuestion,
    SubQuestion,
)
from questions.services import DataImport
from questions.services.expression_scope import (
    sync_expression_dependencies,
    validate_expression_fields,
)

QUESTIONS_XLSX = "./survey_designer/apps/questions/tests/files/questions.xlsx"


def _root_question_payload(question, **overrides):
    return {
        "submodule": [
            str(id_) for id_ in question.submodule.values_list("id", flat=True)
        ],
        "name": question.name,
        "description": question.description,
        "label": question.label,
        "type": question.type,
        "relevant": question.relevant,
        "constraint": question.constraint,
        "calculation": question.calculation,
        "required": question.required,
        "read_only": question.read_only,
        "default": question.default,
        "choice_filter": question.choice_filter,
        "base_question-TOTAL_FORMS": 0,
        "base_question-INITIAL_FORMS": 0,
        "constraint_translations-TOTAL_FORMS": 0,
        "constraint_translations-INITIAL_FORMS": 0,
        "translations-TOTAL_FORMS": 0,
        "translations-INITIAL_FORMS": 0,
        "sub_questions-TOTAL_FORMS": 0,
        "sub_questions-INITIAL_FORMS": 0,
        **overrides,
    }


def _post_root_question(client, question, **overrides):
    url = get_model_admin_base_url(RootQuestion, "_change", [question.id])
    return client.post(url, _root_question_payload(question, **overrides))


def _form_errors(response):
    return response.context["adminform"].form.errors


def test_root_question_admin_rejects_incomplete_expression(
    logged_admin_client, root_question_1, root_question_2
):
    response = _post_root_question(
        logged_admin_client,
        root_question_1,
        relevant=f"${{{root_question_2.name}}} =",
    )

    assert response.status_code == 200
    assert _form_errors(response)["relevant"] == [
        "RootQuestion 'TestQuestion1' field 'relevant' is incomplete after '='; "
        "expected a value, reference, or function."
    ]
    root_question_1.refresh_from_db()
    assert root_question_1.relevant == ""


def test_root_question_admin_rejects_reference_case_mismatch(
    logged_admin_client, root_question_1, root_question_2
):
    response = _post_root_question(
        logged_admin_client,
        root_question_1,
        required=f"${{{root_question_2.name.lower()}}} > 1",
    )

    assert response.status_code == 200
    assert _form_errors(response)["required"] == [
        "RootQuestion 'TestQuestion1' field 'required' references 'testquestion2', "
        "but references are case-sensitive; available exact name: 'TestQuestion2'."
    ]


def test_existing_invalid_record_cannot_be_saved_unchanged(
    logged_admin_client, root_question_1
):
    RootQuestion.objects.filter(id=root_question_1.id).update(
        calculation="round(${TestQuestion1) div 2.0)"
    )
    root_question_1.refresh_from_db()

    response = _post_root_question(logged_admin_client, root_question_1)

    assert response.status_code == 200
    assert "calculation" in _form_errors(response)


def test_root_question_admin_records_every_expression_dependency(
    logged_admin_client, root_question_1, root_question_2
):
    reference = f"${{{root_question_2.name}}}"
    response = _post_root_question(
        logged_admin_client,
        root_question_1,
        required=f"count-selected({reference}) > 0",
        read_only=f"count-selected({reference}) = 0",
        default=f"count-selected({reference}) + 1",
    )

    assert response.status_code == 302
    base_question_id = root_question_2.base_question.id
    for relation in (
        "required_dependencies",
        "read_only_dependencies",
        "default_dependencies",
    ):
        assert list(
            getattr(root_question_1, relation).values_list("id", flat=True)
        ) == [base_question_id]


def test_root_question_admin_rejects_dependency_cycle(
    logged_admin_client, root_question_1, root_question_2
):
    root_question_2.calculation = f"${{{root_question_1.name}}} + 1"
    root_question_2.save(update_fields=["calculation"])

    response = _post_root_question(
        logged_admin_client,
        root_question_1,
        calculation=f"count-selected(${{{root_question_2.name}}})",
    )

    assert response.status_code == 200
    assert _form_errors(response)["calculation"] == [
        "RootQuestion 'TestQuestion1' field 'calculation' creates a dependency "
        "cycle: TestQuestion1 -> TestQuestion2 -> TestQuestion1."
    ]


def test_root_question_admin_rejects_cycle_through_repeat_relevance(
    logged_admin_client, root_question_1, repeat_section_1
):
    repeat_section_1.relevant = f"${{{root_question_1.name}}} > 0"
    repeat_section_1.save(update_fields=["relevant"])

    response = _post_root_question(
        logged_admin_client,
        root_question_1,
        calculation=f"count(${{{repeat_section_1.name}}})",
    )

    assert response.status_code == 200
    assert _form_errors(response)["calculation"] == [
        "RootQuestion 'TestQuestion1' field 'calculation' creates a dependency "
        "cycle: TestQuestion1 -> RepeatSection1 -> TestQuestion1."
    ]


def test_validation_follows_a_rename_through_dependants(
    root_question_1, root_question_2
):
    root_question_2.relevant = f"${{{root_question_1.name}}} > 1"
    root_question_2.save(update_fields=["relevant"])

    issues = validate_expression_fields(
        {"model": "RootQuestion", "id": root_question_1.id, "name": "Renamed"},
        {"relevant": f"count-selected(${{{root_question_2.name}}}) > 0"},
        question_type="integer",
        previous_name=root_question_1.name,
    )

    assert [issue.code for issue in issues] == ["EXPRESSION_DEPENDENCY_CYCLE"]


def test_sub_question_form_rejects_malformed_expression(sub_question_1):
    form_class = modelform_factory(
        SubQuestion,
        form=SubQuestionAdminModelForm,
        fields=("root_question", "suffix", "label", "relevant"),
    )
    form = form_class(
        data={
            "root_question": sub_question_1.root_question_id,
            "suffix": sub_question_1.suffix_id,
            "label": sub_question_1.label,
            "relevant": "${TestQuestion1} and",
        },
        instance=sub_question_1,
    )

    assert form.is_valid() is False
    assert form.errors["relevant"] == [
        f"SubQuestion '{sub_question_1.name}' field 'relevant' is incomplete after "
        "'and'; expected a value, reference, or function."
    ]


def test_repeat_section_form_rejects_count_from_its_own_member(
    repeat_section_1, root_question_2, submodule_1
):
    form_class = modelform_factory(
        RepeatSection,
        form=RepeatSectionAdminModelForm,
        fields=("name", "label", "submodule", "questions", "repeat_count"),
    )
    form = form_class(
        data={
            "name": repeat_section_1.name,
            "label": repeat_section_1.label,
            "submodule": [submodule_1.id],
            "questions": [root_question_2.base_question.id],
            "repeat_count": f"count-selected(${{{root_question_2.name}}})",
        },
        instance=repeat_section_1,
    )

    assert form.is_valid() is False
    assert form.errors["repeat_count"] == [
        "RepeatSection 'RepeatSection1' field 'repeat_count' references "
        "'TestQuestion2', which is inside this repeat; a repeat's repeat_count "
        "cannot depend on its own members."
    ]


def test_repeat_section_form_rejects_non_numeric_count(
    repeat_section_1, root_question_1, root_question_5, submodule_1
):
    RootQuestion.objects.filter(id=root_question_5.id).update(type="date")
    form_class = modelform_factory(
        RepeatSection,
        form=RepeatSectionAdminModelForm,
        fields=("name", "label", "submodule", "questions", "repeat_count"),
    )
    form = form_class(
        data={
            "name": repeat_section_1.name,
            "label": repeat_section_1.label,
            "submodule": [submodule_1.id],
            "questions": [root_question_1.base_question.id],
            "repeat_count": f"${{{root_question_5.name}}}",
        },
        instance=repeat_section_1,
    )

    assert form.is_valid() is False
    assert form.errors["repeat_count"] == [
        "RepeatSection 'RepeatSection1' field 'repeat_count' references "
        "'TestQuestion5' of type 'date', which cannot provide a numeric "
        "repeat count."
    ]


def test_calculation_form_accepts_reference_to_repeat_member(
    root_question_2, repeat_section_1
):
    form_class = modelform_factory(
        Calculation,
        form=CalculationAdminModelForm,
        fields=("name", "label", "calculation"),
    )

    # root_question_2 is a member of a repeat; a calculation has no survey
    # position, so it cannot be "outside" that repeat.
    form = form_class(
        data={
            "name": "MemberChoices",
            "label": "Member choices",
            "calculation": f"count-selected(${{{root_question_2.name}}})",
        }
    )

    assert form.is_valid(), form.errors


def test_calculation_form_rejects_malformed_expression(root_question_1):
    form_class = modelform_factory(
        Calculation,
        form=CalculationAdminModelForm,
        fields=("name", "label", "calculation"),
    )

    form = form_class(
        data={
            "name": "Total",
            "label": "Total",
            "calculation": f"${{{root_question_1.name}}} *",
        }
    )

    assert form.is_valid() is False
    assert form.errors["calculation"] == [
        "Calculation 'Total' field 'calculation' is incomplete after '*'; "
        "expected a value, reference, or function."
    ]


def test_calculate_question_accepts_boolean_alias_calculation():
    assert (
        validate_expression_fields(
            {"model": "RootQuestion", "name": "always_true"},
            {"calculation": "yes"},
            question_type="calculate",
        )
        == []
    )


def test_validation_rejects_calculation_that_reads_itself(root_question_1):
    issues = validate_expression_fields(
        {"model": "RootQuestion", "id": root_question_1.id, "name": "TestQuestion1"},
        {"calculation": ". + 1"},
        question_type="integer",
    )

    assert [issue.code for issue in issues] == ["EXPRESSION_DEPENDENCY_CYCLE"]


def test_dependency_sync_records_repeat_section_references(
    root_question_1, repeat_section_1
):
    root_question_1.calculation = f"count(${{{repeat_section_1.name}}})"
    root_question_1.required = f"count(${{{repeat_section_1.name}}}) > 0"

    sync_expression_dependencies(root_question_1)

    base_question = repeat_section_1.base_question
    assert list(root_question_1.calculation_dependencies.all()) == [base_question]
    assert list(root_question_1.required_dependencies.all()) == [base_question]


def test_dependency_sync_records_exact_references_only(
    root_question_1, root_question_2
):
    root_question_1.relevant = f"${{{root_question_2.name.lower()}}} > 1"
    root_question_1.constraint = f"${{{root_question_2.name}}} > 1"
    root_question_1.calculation = f"${{{root_question_2.name}}} +"

    sync_expression_dependencies(root_question_1)

    assert not root_question_1.relevant_dependencies.exists()
    assert list(root_question_1.constraint_dependencies.all()) == [
        root_question_2.base_question
    ]
    assert not root_question_1.calculation_dependencies.exists()


def test_rename_updates_required_expression(root_question_1, root_question_2):
    root_question_2.required = f"${{{root_question_1.name}}} > 0"
    root_question_2.save(update_fields=["required"])
    sync_expression_dependencies(root_question_2)

    root_question_1.name = "RenamedQuestion"
    root_question_1.save()

    root_question_2.refresh_from_db()
    assert root_question_2.required == "${RenamedQuestion} > 0"


def test_bulk_relevant_view_rejects_malformed_expression(
    logged_admin_client, root_question_1
):
    response = logged_admin_client.post(
        reverse("relevant"),
        {
            "base_questions": [root_question_1.base_question.id],
            "relevant": "${TestQuestion1} and",
        },
    )

    assert response.status_code == 200
    assert response.context["form"].errors["relevant"] == [
        "RootQuestion 'TestQuestion1' field 'relevant' is incomplete after 'and'; "
        "expected a value, reference, or function."
    ]
    root_question_1.refresh_from_db()
    assert root_question_1.relevant == ""


def test_bulk_relevant_view_records_dependencies(
    logged_admin_client, root_question_1, root_question_2
):
    response = logged_admin_client.post(
        reverse("relevant"),
        {
            "base_questions": [root_question_1.base_question.id],
            "relevant": f"count-selected(${{{root_question_2.name}}}) > 0",
        },
    )

    assert response.status_code == 302
    assert list(root_question_1.relevant_dependencies.all()) == [
        root_question_2.base_question
    ]


def test_constraint_builder_joins_conditions_into_a_valid_expression(
    logged_admin_client, root_question_1
):
    response = logged_admin_client.post(
        reverse("constraint"),
        {
            "base_questions": [root_question_1.base_question.id],
            "mode": "form",
            "constraint": "",
            "constraint_message": "",
            "form-TOTAL_FORMS": "2",
            "form-INITIAL_FORMS": "0",
            "form-MIN_NUM_FORMS": "1",
            "form-MAX_NUM_FORMS": "1000",
            "form-0-question": "",
            "form-0-operator": ">",
            "form-0-reference_question": "SELF",
            "form-0-value": "1",
            "form-0-logical_operator": "",
            "form-1-question": "",
            "form-1-operator": "<",
            "form-1-reference_question": "SELF",
            "form-1-value": "5",
            "form-1-logical_operator": "and",
            "translation-TOTAL_FORMS": "0",
            "translation-INITIAL_FORMS": "0",
        },
    )

    assert response.status_code == 302
    root_question_1.refresh_from_db()
    assert root_question_1.constraint == ". > '1' and . < '5'"


def _import_with_survey_cells(admin, cells):
    workbook = load_workbook(QUESTIONS_XLSX)
    survey = workbook["survey"]
    headers = [cell.value for cell in survey[1]]
    for (row, column), value in cells.items():
        survey.cell(row=row, column=headers.index(column) + 1, value=value)
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    data_import = DataImport(
        output, admin, Organization.objects.filter(id=admin.organization.id)
    )
    data_import.process()
    return data_import


def test_bulk_import_preview_reports_field_level_expression_errors(admin):
    data_import = _import_with_survey_cells(
        admin,
        {
            (5, "relevant"): "${FCSStap} >",
            (12, "calculation"): "round(${fcsstap} div 2.0)",
        },
    )

    assert data_import.is_valid() is False
    assert data_import.get_errors()["Questions Spreadsheet"]["errors"] == {
        5: {
            "relevant": [
                "RootQuestion 'FCSStap2' field 'relevant' is incomplete after '>'; "
                "expected a value, reference, or function."
            ]
        },
        12: {
            "calculation": [
                "RootQuestion 'calculation_1' field 'calculation' references "
                "'fcsstap', but references are case-sensitive; available exact "
                "name: 'FCSStap'."
            ]
        },
    }


def test_bulk_import_preview_reports_dependency_cycles_between_rows(admin):
    data_import = _import_with_survey_cells(admin, {(4, "relevant"): "${FCSStap2} > 1"})

    errors = data_import.get_errors()["Questions Spreadsheet"]["errors"]
    assert errors == {
        4: {
            "relevant": [
                "RootQuestion 'FCSStap' field 'relevant' creates a dependency "
                "cycle: FCSStap -> FCSStap2 -> FCSStap."
            ]
        },
        # FCSStap2 also depends on FCSStap through its calculation.
        5: {
            "calculation": [
                "RootQuestion 'FCSStap2' field 'calculation' creates a dependency "
                "cycle: FCSStap2 -> FCSStap -> FCSStap2."
            ]
        },
    }


def test_bulk_import_preview_reports_cycles_through_repeat_relevance(admin):
    data_import = _import_with_survey_cells(
        admin,
        {
            (4, "calculation"): "count(${repeat_test})",
            (9, "relevant"): "${FCSStap} > 0",
        },
    )

    errors = data_import.get_errors()["Questions Spreadsheet"]["errors"]
    assert errors == {
        4: {
            "calculation": [
                "RootQuestion 'FCSStap' field 'calculation' creates a dependency "
                "cycle: FCSStap -> repeat_test -> FCSStap."
            ]
        },
        9: {
            "relevant": [
                "RepeatSection 'repeat_test' field 'relevant' creates a dependency "
                "cycle: repeat_test -> FCSStap -> repeat_test."
            ]
        },
    }


def test_bulk_import_records_every_expression_dependency(admin):
    data_import = _import_with_survey_cells(admin, {(5, "required"): "${FCSStap} > 0"})
    assert data_import.is_valid()

    data_import.create()

    question = RootQuestion.objects.get(name="FCSStap2")
    source = RootQuestion.objects.get(name="FCSStap").base_question
    for relation in (
        "relevant_dependencies",
        "required_dependencies",
        "calculation_dependencies",
    ):
        assert list(getattr(question, relation).all()) == [source]


def test_backfill_migration_records_existing_expression_dependencies(
    root_question_1, root_question_2, sub_question_1, repeat_section_1
):
    migration = importlib.import_module(
        "questions.migrations.0051_backfill_expression_dependencies"
    )
    reference = f"${{{root_question_2.name}}}"
    RootQuestion.objects.filter(id=root_question_1.id).update(
        required=f"count-selected({reference}) > 0",
        read_only="yes",
        default=f"${{{root_question_2.name.lower()}}}",
    )
    SubQuestion.objects.filter(id=sub_question_1.id).update(
        read_only=f"count-selected({reference}) = 0",
        required=f"count(${{{repeat_section_1.name}}}) > 0",
    )

    migration.backfill_expression_dependencies(django_apps, None)

    base_question = root_question_2.base_question
    assert list(root_question_1.required_dependencies.all()) == [base_question]
    assert not root_question_1.read_only_dependencies.exists()
    assert not root_question_1.default_dependencies.exists()
    assert list(sub_question_1.read_only_dependencies.all()) == [base_question]
    assert list(sub_question_1.required_dependencies.all()) == [
        repeat_section_1.base_question
    ]

    root_question_2.name = "RenamedChoices"
    root_question_2.save()

    root_question_1.refresh_from_db()
    assert root_question_1.required == "count-selected(${RenamedChoices}) > 0"


def test_rename_keeps_self_referencing_constraint_valid(
    logged_admin_client, root_question_1
):
    root_question_1.constraint = f"${{{root_question_1.name}}} >= 0"
    root_question_1.save(update_fields=["constraint"])
    sync_expression_dependencies(root_question_1)

    response = _post_root_question(
        logged_admin_client, root_question_1, name="RenamedQuestion"
    )

    assert response.status_code == 302, _form_errors(response)
    root_question_1.refresh_from_db()
    assert root_question_1.name == "RenamedQuestion"
    assert root_question_1.constraint == "${RenamedQuestion} >= 0"
    assert list(root_question_1.constraint_dependencies.all()) == [
        root_question_1.base_question
    ]


def test_bulk_relevant_view_accepts_repeat_section(
    logged_admin_client, repeat_section_1
):
    base_question = BaseQuestion.objects.get(repeat_section=repeat_section_1)

    response = logged_admin_client.post(
        reverse("relevant"),
        {"base_questions": [base_question.id], "relevant": "true()"},
    )

    assert response.status_code == 302
    repeat_section_1.refresh_from_db()
    assert repeat_section_1.relevant == "true()"


def test_bulk_import_records_dependency_on_repeat_from_same_file(admin):
    data_import = _import_with_survey_cells(
        admin, {(5, "required"): "count(${repeat_test}) > 0"}
    )
    assert data_import.is_valid(), data_import.get_errors()

    data_import.create()

    question = RootQuestion.objects.get(name="FCSStap2")
    repeat = BaseQuestion.objects.get(repeat_section__name="repeat_test")
    assert list(question.required_dependencies.all()) == [repeat]


def test_calculation_form_rejects_reference_to_its_own_name():
    form_class = modelform_factory(
        Calculation,
        form=CalculationAdminModelForm,
        fields=("name", "label", "calculation"),
    )

    form = form_class(
        data={
            "name": "calc_template",
            "label": "Template",
            "calculation": "${calc_template} + 1",
        }
    )

    assert form.is_valid() is False
    assert form.errors["calculation"] == [
        "Calculation 'calc_template' field 'calculation' references "
        "'calc_template', but no question with that exact name exists."
    ]
